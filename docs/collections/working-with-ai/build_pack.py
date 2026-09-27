#!/usr/bin/env python3
"""
Build the "Working with AI" pilot pack from pack_source.py. Read-only against the app:
it writes files under this folder (and docs/research/LEGACY_PROVENANCE.csv) and never
touches the database, schema, UI or index.html.

Checks (the build fails if any hard check fails):
  * every row passes the same rules as api/_validate.js#validatePromptRow
  * ids and titles do not collide with the committed library (web/seed/prompts.json)
  * Synottic's own framework detector rates each prompt at or above its intended level
  * the audit's safety detectors find nothing (jailbreak, NSFW, malicious, deception,
    synthetic identity, tool/model names)
  * overlap with the nearest existing Synottic prompt is reported (warning only)

Usage (from the repo root):
  python3 docs/collections/working-with-ai/build_pack.py --pc <prompts.chat checkout>
Requires: scikit-learn, openpyxl, node.
"""
import argparse
import csv
import json
import os
import re
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", "..", ".."))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(REPO, "docs", "research", "tools"))
csv.field_size_limit(10**9)

import pack_source as S  # noqa: E402
import prompts_chat_audit as AUDIT  # noqa: E402

# Mirrors api/_validate.js
CATEGORIES = [
    "AI & Prompt Engineering", "Book & Ebook Writing", "Business Strategy", "Career Growth",
    "Coaching & Self-Development", "Coding & Tech", "Communication & Leadership",
    "Content Writing & Copywriting", "Customer Support", "E-Commerce", "Education & Learning",
    "Email Marketing", "Finance & Accounting", "General", "Health & Fitness", "HR & Recruiting",
    "Image & Design", "Legal & Compliance", "Marketing & Branding", "Presentation & Slides",
    "Product Management", "Productivity & Automation", "Research & Data Analysis",
    "SEO & Analytics", "Sales & Lead Generation", "Social Media", "Spirituality & Wellness",
    "UX/UI Design",
]
DIFFICULTIES = ["Beginner", "Intermediate", "Advanced"]
LIFECYCLE = "Curated"


def variables(text):
    seen = []
    for v in re.findall(r"\[([A-Z][A-Z0-9_]{1,40})\]", text):
        if v not in seen:
            seen.append(v)
    return seen


def row_for(p, category, extra_tags):
    return {
        "id": p["id"], "title": p["title"], "category": category, "role": p["role"],
        "useCase": f'{p["title"]}. {p["description"]}', "description": p["description"],
        "outcome": p["outcome"], "originalPrompt": p["prompt"].strip(),
        "promptType": p["promptType"], "difficulty": p["difficulty"],
        "variables": variables(p["prompt"]), "tags": p["tags"] + extra_tags, "lifecycle": LIFECYCLE,
    }


def validate(r, errors):
    t = r["title"].strip()
    if not 3 <= len(t) <= 300:
        errors.append(f"{r['id']}: title length")
    if r["category"] not in CATEGORIES:
        errors.append(f"{r['id']}: invalid category {r['category']}")
    if not 10 <= len(r["originalPrompt"]) <= 20000:
        errors.append(f"{r['id']}: prompt length")
    if r["difficulty"] not in DIFFICULTIES:
        errors.append(f"{r['id']}: invalid difficulty")
    for k, cap in (("role", 160), ("useCase", 400), ("description", 1000), ("promptType", 60), ("outcome", 400)):
        if len(r[k]) > cap:
            errors.append(f"{r['id']}: {k} longer than {cap} (would be truncated)")
    if len(r["variables"]) > 60 or len(r["tags"]) > 40:
        errors.append(f"{r['id']}: too many variables/tags")
    if "—" in r["originalPrompt"]:
        errors.append(f"{r['id']}: em dash in prompt body (house style)")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pc", required=True, help="path to a checkout of f/prompts.chat at the audited commit")
    a = ap.parse_args()

    seed = json.load(open(os.path.join(REPO, "web", "seed", "prompts.json"), encoding="utf-8"))
    seed_by_id = {r["id"]: r for r in seed}
    pc_rows = list(csv.DictReader(open(os.path.join(a.pc, "prompts.csv"), encoding="utf-8")))
    mapping = {r["source_prompt_id"]: r for r in csv.DictReader(open(os.path.join(REPO, "docs", "research", "PROMPTS_CHAT_MAPPING.csv"), encoding="utf-8"))}

    errors, warnings = [], []
    pilot = [row_for(p, S.CATEGORY, [S.TAG, f"L{p['level']}"]) for p in S.PROMPTS]
    cleanup = []
    for c in S.CLEANUP:
        base = seed_by_id.get(c["id"])
        if not base:
            errors.append(f"{c['id']}: not in the committed library")
            continue
        if base["title"] != c["title"] or base["category"] != c["category"]:
            errors.append(f"{c['id']}: title/category must match the live record so import updates it in place")
        cleanup.append(row_for(dict(c, sources=[], level=None), c["category"], []))

    # ── rule checks
    for r in pilot + cleanup:
        validate(r, errors)
    ids = [r["id"] for r in pilot]
    if len(set(ids)) != len(ids):
        errors.append("duplicate ids in pilot")
    norm = lambda s: re.sub(r"\s+", " ", s.strip().lower())
    seen_titles = {(norm(r["title"]), r["category"]): r["id"] for r in seed}
    for r in pilot:
        if r["id"] in seed_by_id:
            errors.append(f"{r['id']}: id already exists in library")
        hit = seen_titles.get((norm(r["title"]), r["category"]))
        if hit:
            errors.append(f"{r['id']}: title already used by {hit} in {r['category']} (import would overwrite it)")

    # ── framework level (Synottic's own detector)
    fw_in = json.dumps([{"id": r["id"], "originalPrompt": r["originalPrompt"]} for r in pilot + cleanup])
    fw = json.loads(subprocess.run(["node", os.path.join(HERE, "check_framework.mjs"), os.path.join(REPO, "src", "part_framework.js")],
                                   input=fw_in, capture_output=True, text=True, check=True).stdout)
    fw_by_id = {x["id"]: x for x in fw}
    intended = {p["id"]: p["level"] for p in S.PROMPTS}
    intended.update({c["id"]: (3 if "Goal:" in c["prompt"] else 2) for c in S.CLEANUP})
    for rid, lvl in intended.items():
        got = fw_by_id[rid]["level"]
        if got is None or got < lvl:
            errors.append(f"{rid}: framework detector says L{got}, intended L{lvl}")

    # ── safety scan (same detectors as the Phase 1 audit)
    for r in pilot + cleanup:
        t, body = r["title"], r["originalPrompt"]
        hits = []
        if AUDIT.is_jailbreak(t, body): hits.append("jailbreak")
        if AUDIT.flagged(AUDIT.NSFW, t + " " + body, 60): hits.append("nsfw")
        if AUDIT.flagged(AUDIT.MALICIOUS, t + " " + body): hits.append("malicious")
        if AUDIT.flagged(AUDIT.SYNTHETIC_IDENTITY, t + " " + body): hits.append("synthetic identity")
        if AUDIT.DECEPTIVE.search(t + " " + body): hits.append("deception")
        if AUDIT.TOOLISH.search(body): hits.append("tool/model name: " + AUDIT.TOOLISH.search(body).group(0))
        if hits:
            errors.append(f"{r['id']}: safety scan hit {hits}")

    # ── overlap with the existing library (warning only; pilot rows only)
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.metrics.pairwise import cosine_similarity

    def doc(title, desc, prompt):
        task = re.search(r"^Task:\s*(.*)$", prompt, re.M)
        return " ".join([title, title, desc, task.group(1) if task else ""])
    lib = [r for r in seed if r["id"] not in {c["id"] for c in S.CLEANUP}]
    lib_docs = [doc(r["title"], r.get("description") or "", r["originalPrompt"] or "") for r in lib]
    p_docs = [doc(r["title"], r["description"], r["originalPrompt"]) for r in pilot]
    vec = TfidfVectorizer(ngram_range=(1, 2), stop_words="english", sublinear_tf=True).fit(lib_docs + p_docs)
    sim = cosine_similarity(vec.transform(p_docs), vec.transform(lib_docs))
    overlap = {}
    for i, r in enumerate(pilot):
        j = int(sim[i].argmax())
        overlap[r["id"]] = (lib[j]["id"], lib[j]["title"], float(sim[i, j]))
        if sim[i, j] >= 0.45:
            warnings.append(f"{r['id']}: close to {lib[j]['id']} '{lib[j]['title']}' ({sim[i, j]:.2f}); confirm it is a different job")

    # the level tag follows what the app will actually detect
    for r in pilot:
        r["tags"] = [t for t in r["tags"] if not re.fullmatch(r"L[123]", t)] + [f"L{fw_by_id[r['id']]['level']}"]

    if errors:
        print("BUILD FAILED:\n  " + "\n  ".join(errors))
        sys.exit(1)

    # ── outputs
    os.makedirs(os.path.join(HERE, "import"), exist_ok=True)
    json.dump(cleanup, open(os.path.join(HERE, "import", "01_cleanup.json"), "w", encoding="utf-8"), indent=2, ensure_ascii=False)
    json.dump(pilot, open(os.path.join(HERE, "import", "02_working_with_ai.json"), "w", encoding="utf-8"), indent=2, ensure_ascii=False)

    def src_info(sid):
        i = int(sid.split("-")[1])
        pr = pc_rows[i - 1]
        return pr, mapping.get(sid, {})

    with open(os.path.join(HERE, "provenance.csv"), "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["synottic_id", "synottic_title", "source", "source_ids", "source_titles", "source_locator",
                    "source_commit", "source_url", "licence", "contributors", "transformation", "what_changed", "shown_in_product"])
        for p in S.PROMPTS:
            infos = [src_info(s) for s in p["sources"]]
            w.writerow([p["id"], p["title"], "prompts.chat", " | ".join(p["sources"]),
                        " | ".join(pr["act"] for pr, _ in infos),
                        " | ".join(f"prompts.csv row {int(s.split('-')[1])}" for s in p["sources"]),
                        S.SOURCE_COMMIT, "https://github.com/f/prompts.chat", "CC0-1.0",
                        " | ".join("@" + pr["contributor"] for pr, _ in infos), "rewritten", p["changes"], "no"])

    # legacy records that already hold prompts.chat text (Phase 1 finding)
    legacy = []
    blob = {r["id"]: re.sub(r"\s+", " ", (r.get("legacyPrompt") or "").lower()) for r in seed if r.get("legacyPrompt")}
    for i, pr in enumerate(pc_rows, 1):
        words = re.sub(r"\s+", " ", pr["prompt"].lower()).split()
        if len(words) < 14:
            continue
        opening = " ".join(words[:12])
        for rid, text in blob.items():
            if opening in text:
                legacy.append((rid, seed_by_id[rid]["title"], f"pc-{i:04d}", pr["act"], pr["contributor"]))
    legacy.sort()
    with open(os.path.join(REPO, "docs", "research", "LEGACY_PROVENANCE.csv"), "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["synottic_id", "synottic_title", "source", "source_id", "source_title", "source_locator",
                    "source_commit", "licence", "contributor", "where_found", "shown_in_product"])
        for rid, title, sid, stitle, contrib in legacy:
            w.writerow([rid, title, "prompts.chat", sid, stitle, f"prompts.csv row {int(sid.split('-')[1])}", S.SOURCE_COMMIT,
                        "CC0-1.0", "@" + contrib, "legacyPrompt (pre-rewrite text; not displayed)", "no"])

    # collection settings
    with open(os.path.join(HERE, "collection.md"), "w", encoding="utf-8") as fh:
        fh.write(f"""# Admin Collection: {S.COLLECTION_NAME}

Enter these in **Admin console → Collections → New collection** after `import/02_working_with_ai.json` has been committed.

| Field | Value |
|---|---|
| Name | {S.COLLECTION_NAME} |
| Organisation | the organisation this collection is for (one collection per organisation) |
| Categories | leave empty (the collection is the 20 pinned prompts only) |
| Programs | leave empty |
| Joining defaults | Function: `general` · AI level: `beginner` (or match the organisation) |

**Pinned prompt IDs** (paste into "Pinned prompt IDs"):

```
{chr(10).join(ids)}
```

The console shows "20/20 pinned IDs match" once the imported prompts have loaded. If it shows fewer, reload the console so it picks up the newly imported prompts.

Then use **Generate code** on the collection to create the access code learners redeem.

## Suggested re-filing (not applied)

The clean-up rewrites keep each prompt's current category so nobody's library scope changes unexpectedly. Their content suggests better homes, for a later re-filing pass:

| Prompt | Current category | Suggested |
|---|---|---|
| `lib-336` Translate resume bullets across fields | Coding & Tech | Career Growth |
| `lib-437` Avoid plagiarism in written content | Coding & Tech | Content Writing & Copywriting |
| `lib-2178` Make AI text read naturally | Legal & Compliance | Content Writing & Copywriting |
""")

    # review workbook
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F3A5F")
    wrap = Alignment(wrap_text=True, vertical="top")

    def sheet(ws, headers, rows, widths):
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill, c.alignment = head, fill, wrap
        for r in rows:
            ws.append(r)
        for i, wdt in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = wdt
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = wrap
        ws.freeze_panes = "C2"

    ws = wb.active
    ws.title = "Pilot prompts"
    rows = []
    for n, p in enumerate(S.PROMPTS, 1):
        r = next(x for x in pilot if x["id"] == p["id"])
        infos = [src_info(s) for s in p["sources"]]
        src_text = "\n\n---\n\n".join(" ".join(pr["prompt"].split())[:1500] + (" …" if len(pr["prompt"]) > 1500 else "") for pr, _ in infos)
        o = overlap[p["id"]]
        rows.append([n, p["id"], p["title"], f"L{fw_by_id[p['id']]['level']}", p["role"],
                     p["promptType"], p["difficulty"], r["originalPrompt"], ", ".join(r["variables"]),
                     " | ".join(p["sources"]), " | ".join(pr["act"] for pr, _ in infos), src_text, p["changes"],
                     f"{o[0]} '{o[1]}' ({o[2]:.2f})", "", ""])
    sheet(ws, ["#", "Synottic ID", "Title", "Framework level", "Role", "Type", "Difficulty", "Rewritten prompt (Synottic)",
               "Variables", "Source ID", "Source title", "Source prompt (prompts.chat, excerpt)", "What changed",
               "Nearest existing Synottic prompt (overlap)", "Approve? (Y / N / Edit)", "Reviewer notes"],
          rows, [4, 13, 30, 12, 20, 12, 11, 70, 24, 11, 28, 60, 40, 32, 12, 30])

    ws2 = wb.create_sheet("Clean-up")
    rows2 = []
    for c in S.CLEANUP:
        base = seed_by_id[c["id"]]
        rows2.append([c["id"], c["title"], c["why"], base["originalPrompt"], c["prompt"].strip(), "", ""])
    sheet(ws2, ["ID", "Title", "Why", "Current live prompt", "Replacement", "Approve? (Y / N / Edit)", "Reviewer notes"],
          rows2, [11, 30, 40, 70, 70, 12, 30])

    ws3 = wb.create_sheet("Legacy provenance")
    sheet(ws3, ["Synottic ID", "Synottic title", "Source ID", "Source title", "Contributor"],
          [[a1, b, c1, d, "@" + e] for a1, b, c1, d, e in legacy], [12, 40, 11, 36, 24])

    ws4 = wb.create_sheet("How to load")
    steps = [
        ["1", "Review", "Mark each row in 'Pilot prompts' and 'Clean-up' as Y, N or Edit. Only approved rows are loaded."],
        ["2", "Clean-up import", "Admin console → Prompts → Import → import/01_cleanup.json. Preview must show 3 × update, 0 new. Commit."],
        ["3", "Pilot import", "Admin console → Prompts → Import → import/02_working_with_ai.json. Preview must show 20 × new. Commit."],
        ["4", "Collection", "Admin console → Collections → New. Follow collection.md (name, organisation, 20 pinned IDs). Save, then Generate code."],
        ["5", "Check", "Redeem the code with a test learner account and confirm exactly 20 prompts appear."],
        ["6", "Measure", "After 2 to 4 weeks, compare opens, copies and saves for these 20 against other AI & Prompt Engineering prompts."],
    ]
    sheet(ws4, ["Step", "What", "How"], steps, [6, 18, 110])
    ws4.freeze_panes = "A2"
    wb.save(os.path.join(HERE, "review.xlsx"))

    print(f"OK: {len(pilot)} pilot rows, {len(cleanup)} clean-up rows, {len(legacy)} legacy provenance links")
    for rid in intended:
        print(f"  {rid}: intended L{intended[rid]}, detector L{fw_by_id[rid]['level']}, overlap {overlap.get(rid, ('', '', 0))[2]:.2f} {overlap.get(rid, ('', ''))[0]}")
    for wmsg in warnings:
        print("WARN", wmsg)


if __name__ == "__main__":
    main()
